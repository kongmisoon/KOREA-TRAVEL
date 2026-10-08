"""가짜(개인정보 없음) HIRA 진료내역 Excel 생성기. 가정한 컬럼 구성으로 만든 더미 데이터.

기대 JSON 재생성:  python backend/services/normalization/tests/fixtures/make_sample.py
"""
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

HEADER = ["진료일자", "약품코드", "약품명", "주성분명", "1회\n투약량", "1일 투여횟수", "총 투약일수", "요양기관명"]
D = dt.datetime


def build_sample_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "진료내역"
    ws["A1"] = "진료내역 조회결과 (더미)"
    ws.merge_cells("A1:H1")
    ws.append([])  # 빈 행
    ws.append(HEADER)  # 3행
    rows = [
        [D(2024, 3, 5), "200800001", "하나펜정500mg", "아세트아미노펜", 1, 3, 3, "가나다의원"],            # 4  정상 CODE_EXACT
        [D(2024, 3, 5), None, "하나펜정650mg", None, 1, 2, 3, None],                                       # 5  기관 병합, 함량으로 구분 NAME_EXACT
        ["2024.03.06", "012345678", "셋트린정10mg", None, 1, 1, 7, "라마바약국"],                          # 6  앞자리 0 코드
        ["2024-13-45", 12345678, "둘렉신캡슐250mg", None, 1, 3, 5, "라마바약국"],                          # 7  이상 날짜 + 숫자 코드
        [],                                                                                                # 8  빈 행
        [None, None, "하나펜정", None, 1, 3, 3, "가나다의원"],                                              # 9  빈 날짜/코드, 중복 후보
        ["20240307", None, "베타글루칸프로로정", None, 1, 1, 14, "가나다의원"],                             # 10 유사도 단일
        [D(2024, 3, 8), None, "엡실론콤비플러스액티브", None, 1, 1, 14, "가나다의원"],                      # 11 유사도 복수
        [D(2024, 3, 8), None, "존재안함정", None, 1, 1, 5, "가나다의원"],                                  # 12 매칭 실패
        HEADER,                                                                                            # 13 반복 헤더
        [D(2024, 3, 9), None, None, None, 1, 1, 5, "가나다의원"],                                          # 14 약품명 없음 -> FAILED
        [D(2024, 3, 10), "H2020000001", "하루비타민C 1000mg", None, 0, 1, "30일", "가나다의원"],           # 15 건기식, 복용량 0
        [D(2024, 3, 10), None, "성분없는정", None, "abc", 1, 5, "가나다의원"],                              # 16 공식 성분 없음, 숫자 오류
    ]
    for r in rows:
        ws.append(r)
    ws.merge_cells("H4:H5")  # 기관명 세로 병합
    wb.save(path)


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[4]
    sys.path.insert(0, str(root.parent))
    from backend.services.normalization.service import normalize_file

    out = Path(__file__).parent
    xlsx = out / "_sample_tmp.xlsx"
    build_sample_xlsx(xlsx)
    results, fw, fe = normalize_file(xlsx)
    xlsx.unlink()
    (out / "expected_sample.json").write_text(
        json.dumps({"file_warnings": fw, "file_errors": fe, "results": [r.to_dict() for r in results]},
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
