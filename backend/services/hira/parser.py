"""HIRA 진료내역 Excel -> MedicationRecord (규칙 기반, LLM 미사용, 결정적)."""
from __future__ import annotations

import datetime as dt
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

import openpyxl
import yaml

from backend.services.normalization.models import (
    Dosage, Ingredient, MedicationRecord, RawNormalized, Strength,
)
from backend.services.normalization.name_utils import clean_text, parse_strength

log = logging.getLogger(__name__)
DEFAULT_MAPPING = Path(__file__).with_name("column_mapping.yaml")


@dataclass
class ParsedRow:
    sheet: str
    row_number: int
    record: Optional[MedicationRecord] = None  # 실패하면 None
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


@dataclass
class ParseOutput:
    rows: list[ParsedRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def load_mapping(path: Optional[Path] = None) -> dict:
    with open(path or DEFAULT_MAPPING, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _header_key(value: Any) -> str:
    s = re.sub(r"\s+", "", clean_text(str(value))).lower() if value is not None else ""
    return s


def _header_variants(value: Any) -> set[str]:
    key = _header_key(value)
    return {key, re.sub(r"\(.*?\)", "", key)} - {""}


def _cell_text(value: Any) -> Optional[str]:
    """원본 보존용 문자열. 빈 셀/공백만 있으면 None."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    s = str(value).strip()
    return s or None


def _load_grid(ws) -> list[list[Any]]:
    """시트를 2차원 값 배열로. 병합 셀은 병합 영역 전체를 좌상단 값으로 채운다."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
            for c in range(rng.min_col - 1, min(rng.max_col, len(grid[r]))):
                grid[r][c] = top
    return grid


def _find_header(grid: list[list[Any]], aliases: dict[str, set[str]], scan: int, minimum: int):
    """별칭이 가장 많이 매칭되는 행을 헤더로. (행 index, {field: col index})"""
    best: tuple[int, int, dict] = (0, -1, {})
    for r, row in enumerate(grid[:scan]):
        cols: dict[str, int] = {}
        for c, v in enumerate(row):
            for fld, names in aliases.items():
                if fld not in cols and _header_variants(v) & names:
                    cols[fld] = c
                    break
        if len(cols) > best[0]:
            best = (len(cols), r, cols)
    n, r, cols = best
    if n < minimum or "product_name" not in cols:
        return None
    return r, cols


def parse_date(value: Any) -> tuple[Optional[str], Optional[str]]:
    """-> (YYYY-MM-DD | None, 경고 사유 | None). 날짜 셀이 비어 있으면 (None, None)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    try:
        if isinstance(value, dt.datetime):
            return value.date().isoformat(), None
        if isinstance(value, dt.date):
            return value.isoformat(), None
        s = clean_text(str(value))
        if isinstance(value, float) and value.is_integer():
            s = str(int(value))
        m = (re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s)
             or re.fullmatch(r"(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})\.?", s)
             or re.fullmatch(r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일?", s))
        if not m:
            return None, f"DATE_UNPARSEABLE: '{value}'"
        return dt.date(int(m[1]), int(m[2]), int(m[3])).isoformat(), None
    except ValueError:
        return None, f"DATE_INVALID: '{value}'"


def parse_number(value: Any) -> tuple[Optional[float], Optional[str]]:
    """'1', '0.5', '1/2', '3정', '1,000' -> float. 비어 있으면 (None, None), 해석 불가면 (None, 사유)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, bool):
        return None, f"NUMBER_UNPARSEABLE: '{value}'"
    if isinstance(value, (int, float)):
        return float(value), None
    s = clean_text(str(value)).replace(" ", "")
    if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
        s = s.replace(",", "")
    m = re.fullmatch(r"(\d+(?:\.\d+)?|\d+/\d+)[가-힣a-zA-Z]*", s)
    if not m:
        return None, f"NUMBER_UNPARSEABLE: '{value}'"
    t = m.group(1)
    if "/" in t:
        a, b = t.split("/")
        return (None, f"NUMBER_UNPARSEABLE: '{value}'") if float(b) == 0 else (float(a) / float(b), None)
    return float(t), None


def _parse_code(value: Any, warnings: list[str]) -> tuple[Optional[str], Optional[str]]:
    """drug_code는 항상 String. 숫자 셀은 앞자리 0이 이미 사라졌을 수 있어 경고만 남기고 패딩하지 않는다."""
    if value is None:
        return None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        code = str(int(value)) if float(value).is_integer() else str(value)
        warnings.append(f"CODE_NUMERIC_CELL: '{code}' 숫자 셀이라 앞자리 0이 유실됐을 수 있음 (보정하지 않음)")
        return code, code
    s = str(value).strip()
    return (s or None), (s or None)


def _parse_row(sheet: str, rno: int, row: list[Any], cols: dict[str, int], headers: list[str]) -> ParsedRow:
    out = ParsedRow(sheet=sheet, row_number=rno)
    get = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
    w = out.warnings

    raw_name = _cell_text(get("product_name"))
    if raw_name is None:
        out.errors.append("PRODUCT_NAME_MISSING: 약품명이 비어 있어 행을 해석할 수 없음")
        return out

    rec = MedicationRecord(sheet=sheet, row_number=rno)
    rec.raw_row = {h: _cell_text(v) for h, v in zip(headers, row) if h}
    rec.product_name = RawNormalized(raw=raw_name, normalized=clean_text(raw_name))

    rec.date_raw = _cell_text(get("date"))
    rec.date, why = parse_date(get("date"))
    if why:
        w.append(why)
    elif rec.date is None and "date" in cols:
        w.append("DATE_MISSING: 날짜 없음 -> null")

    rec.drug_code, rec.drug_code_raw = _parse_code(get("drug_code"), w)
    if rec.drug_code is None and "drug_code" in cols:
        w.append("CODE_MISSING: 약품코드 없음 -> null")

    # 함량: 컬럼 우선, 없으면 제품명 안의 표기에서 (규칙 기반 추출)
    raw_strength = _cell_text(get("strength"))
    if raw_strength:
        v, u = parse_strength(raw_strength)
        rec.strength = Strength(raw=raw_strength, value=v, unit=u, source="COLUMN")
        if v is None:
            w.append(f"STRENGTH_UNPARSEABLE: '{raw_strength}'")
    else:
        v, u = parse_strength(raw_name)
        if v is not None:
            m = re.search(r"\d+(?:\.\d+)?\s*\S+", raw_name)
            rec.strength = Strength(raw=m.group(0) if m else None, value=v, unit=u, source="PRODUCT_NAME")
        else:
            w.append("STRENGTH_MISSING: 함량 정보 없음 -> null")

    raw_ing = _cell_text(get("ingredient"))
    if raw_ing:
        for part in re.split(r"[,;/+]", raw_ing):
            if part.strip():
                rec.ingredients.append(Ingredient(name=clean_text(part), source="HIRA"))

    d = Dosage(raw={k: _cell_text(get(k)) for k in ("dose_per_time", "times_per_day", "total_days")})
    d.dose_per_time, why = parse_number(get("dose_per_time"))
    if why:
        w.append(why + " (dose_per_time)")
    d.times_per_day, why = parse_number(get("times_per_day"))
    if why:
        w.append(why + " (times_per_day)")
    days, why = parse_number(get("total_days"))
    if why:
        w.append(why + " (total_days)")
    if days is not None and not days.is_integer():
        w.append(f"TOTAL_DAYS_NOT_INTEGER: '{days}' -> null")
        days = None
    d.total_days = int(days) if days is not None else None
    rec.dosage = d

    rec.institution = _cell_text(get("institution"))
    out.record = rec
    return out


def parse_hira_excel(path: str | Path, mapping_path: Optional[Path] = None) -> ParseOutput:
    """Excel 전체를 파싱한다. 한 행이 실패해도 계속 진행하고 그 행만 errors에 기록한다."""
    result = ParseOutput()
    mapping = load_mapping(mapping_path)
    aliases = {f: {_header_key(a) for a in names} for f, names in mapping["fields"].items()}
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # 파일 자체를 열 수 없음
        result.errors.append(f"FILE_UNREADABLE: {type(e).__name__}: {e}")
        return result

    found = False
    for ws in wb.worksheets:
        try:
            grid = _load_grid(ws)
            hit = _find_header(grid, aliases, mapping["header_scan_rows"], mapping["min_header_matches"])
            if hit is None:
                result.warnings.append(f"SHEET_SKIPPED: '{ws.title}' 에서 헤더 행을 찾지 못함")
                continue
            found = True
            hr, cols = hit
            headers = [_cell_text(v) or "" for v in grid[hr]]
            for i, row in enumerate(grid[hr + 1:], start=hr + 2):
                if all(_cell_text(v) is None for v in row):
                    continue  # 빈 행
                if _is_repeated_header(row, grid[hr]):
                    continue  # 페이지마다 반복되는 헤더
                try:
                    result.rows.append(_parse_row(ws.title, i, row, cols, headers))
                except Exception as e:  # 한 행 실패가 전체를 멈추지 않게
                    log.exception("row failed sheet=%s row=%s", ws.title, i)
                    result.rows.append(ParsedRow(ws.title, i, errors=[f"ROW_EXCEPTION: {type(e).__name__}: {e}"]))
        except Exception as e:
            result.warnings.append(f"SHEET_FAILED: '{ws.title}': {type(e).__name__}: {e}")
    if not found:
        result.errors.append("HEADER_NOT_FOUND: 어느 시트에서도 헤더 행을 찾지 못함 (column_mapping.yaml 확인)")
    return result


def _is_repeated_header(row: list[Any], header_row: list[Any]) -> bool:
    a = [_header_key(v) for v in row if _header_key(v)]
    b = [_header_key(v) for v in header_row if _header_key(v)]
    return bool(a) and a == b
